from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import openpyxl
from openpyxl import load_workbook
from datetime import datetime, timedelta
import shutil, re, os, tempfile, zipfile, json, io
import pymupdf
import requests
import gspread
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
import pytz

app = Flask(__name__)
CORS(app)

TEMPLATE_PATH      = os.path.join(os.path.dirname(__file__), 'template.xlsx')
TEMPLATE_BU1_PATH  = os.path.join(os.path.dirname(__file__), 'template_bu1.pdf')
TEMPLATE_NAME_FILE = os.path.join(os.path.dirname(__file__), 'template_name.txt')
WEEKDAY_MAP = ['一','二','三','四','五','六','日']
MAX_PER_SHEET = 10
RENDER_SERVICE_ID = 'srv-d7mecdb7uimc73crjh6g'
SCOPES = ['https://spreadsheets.google.com/feeds','https://www.googleapis.com/auth/drive']

# Google Drive 範本 File ID（從 Drive 連結取得）
DRIVE_TEMPLATE_FILE_ID = '1uPD8cgYuOoD95WZBgoyXjSg-54G1YBNO'

# ── 台灣時區 ────────────────────────────────────────────────
TW_TZ = pytz.timezone('Asia/Taipei')

def tw_now():
    return datetime.now(TW_TZ)

# ── Google Drive 工具函式 ────────────────────────────────────
def get_drive_service():
    """取得 Google Drive API 服務"""
    try:
        creds_json = os.environ.get('GOOGLE_SERVICE_ACCOUNT', '')
        if not creds_json:
            return None
        creds_dict = json.loads(creds_json)
        creds = Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
        return build('drive', 'v3', credentials=creds)
    except Exception as e:
        print(f'Drive 服務初始化失敗: {e}')
        return None

def download_template_from_drive():
    """從 Google Drive 下載最新範本到本地"""
    try:
        service = get_drive_service()
        if not service:
            print('Drive 服務不可用，跳過範本下載')
            return False
        # 以 xlsx 格式匯出（因為 Drive 上是 Google Sheets 格式）
        request_ = service.files().export_media(
            fileId=DRIVE_TEMPLATE_FILE_ID,
            mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        buf = io.BytesIO()
        downloader = MediaIoBaseDownload(buf, request_)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buf.seek(0)
        with open(TEMPLATE_PATH, 'wb') as f:
            f.write(buf.read())
        print(f'✅ 範本已從 Drive 下載更新：{TEMPLATE_PATH}')
        return True
    except Exception as e:
        print(f'從 Drive 下載範本失敗: {e}')
        return False

def upload_template_to_drive():
    """把本地範本上傳（更新）到 Google Drive"""
    try:
        service = get_drive_service()
        if not service:
            return False
        with open(TEMPLATE_PATH, 'rb') as f:
            media = MediaIoBaseUpload(
                io.BytesIO(f.read()),
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                resumable=False
            )
        service.files().update(
            fileId=DRIVE_TEMPLATE_FILE_ID,
            media_body=media
        ).execute()
        print('✅ 範本已上傳到 Drive')
        return True
    except Exception as e:
        print(f'上傳範本到 Drive 失敗: {e}')
        return False

# ── 啟動時從 Drive 下載最新範本 ──────────────────────────────
def init_template():
    """應用程式啟動時，從 Drive 下載最新範本"""
    print('=== 啟動：從 Drive 同步範本 ===')
    download_template_from_drive()

# 在 Flask app 啟動時執行
with app.app_context():
    init_template()


def get_sheet():
    try:
        creds_json = os.environ.get('GOOGLE_SERVICE_ACCOUNT','')
        sheet_id   = os.environ.get('GOOGLE_SHEET_ID','')
        if not creds_json or not sheet_id:
            return None
        import json
        creds_dict = json.loads(creds_json)
        creds  = Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
        client = gspread.authorize(creds)
        sheet  = client.open_by_key(sheet_id)
        ws = sheet.sheet1
        if not ws.row_values(1):
            ws.append_row(['操作時間','操作者','課別','版本一檔名','版本二檔名',
                          '依據版本','總異動店數','異動通知書','補1通知書',
                          '店號','店名','原訂日期','原訂午別','異動後日期','異動後午別','文件類型'])
        return ws
    except Exception as e:
        print(f'Google Sheets 連線失敗: {e}')
        return None

def log_to_sheet(ws, operator, dept, f1name, f2name, version_full, 異動通知, 補1通知):
    try:
        if not ws:
            return
        now = tw_now().strftime('%Y/%m/%d %H:%M')
        total = len(異動通知) + len(補1通知)
        rows = []
        for d in 異動通知:
            rows.append([now, operator, dept, f1name, f2name, version_full,
                        total, len(異動通知), len(補1通知),
                        d['店號'], d['店名'], d['原訂日'], d['原訂午'],
                        d['異動日'], d['異動午'], '異動通知書'])
        for d in 補1通知:
            rows.append([now, operator, dept, f1name, f2name, version_full,
                        total, len(異動通知), len(補1通知),
                        d['店號'], d['店名'], d['原訂日'], d['原訂午'],
                        d['異動日'], d['異動午'], '補1通知書'])
        if rows:
            ws.append_rows(rows)
    except Exception as e:
        print(f'寫入 Google Sheets 失敗: {e}')

# ── Google Sheets 設定管理 ───────────────────────────────────
def get_settings_sheet():
    try:
        creds_json = os.environ.get('GOOGLE_SERVICE_ACCOUNT','')
        sheet_id   = os.environ.get('GOOGLE_SHEET_ID','')
        if not creds_json or not sheet_id:
            return None
        import json
        creds_dict = json.loads(creds_json)
        creds  = Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
        client = gspread.authorize(creds)
        sheet  = client.open_by_key(sheet_id)
        return sheet.worksheet('系統設定')
    except Exception as e:
        print(f'系統設定 Sheets 連線失敗: {e}')
        return None

def get_setting(key, default=''):
    try:
        ws = get_settings_sheet()
        if not ws: return default
        rows = ws.get_all_values()
        for row in rows[1:]:
            if row and row[0] == key:
                return row[1] if len(row) > 1 else default
        return default
    except:
        return default

def set_setting(key, value):
    try:
        ws = get_settings_sheet()
        if not ws: return False
        rows = ws.get_all_values()
        for i, row in enumerate(rows):
            if row and row[0] == key:
                ws.update_cell(i+1, 2, value)
                return True
        ws.append_row([key, value])
        return True
    except Exception as e:
        print(f'set_setting 失敗: {e}')
        return False

def get_admin_password():
    return get_setting('ADMIN_PASSWORD', os.environ.get('ADMIN_PASSWORD', ''))

def get_template_name():
    return get_setting('TEMPLATE_NAME', 'template.xlsx')

def save_template_name(name):
    set_setting('TEMPLATE_NAME', name)

def get_notice_base_name():
    name = get_template_name()
    base = os.path.splitext(name)[0]
    return base

def get_notice_filename(first_date_str):
    """
    依據班表第一筆日期產生異動通知書檔名前綴。
    格式：(民國年+月份)盤點行程異動通知書-
    例如：first_date_str = '2026/04/02' → '(11504)盤點行程異動通知書-'
    民國年 = 西元年 - 1911，2026 → 115，2027 → 116
    月份取自班表第一筆資料的月份
    """
    try:
        parts = first_date_str.split('/')
        year_roc = int(parts[0]) - 1911   # 2026 - 1911 = 115
        month = int(parts[1])
        return f'({year_roc}{month:02d})盤點行程異動通知書-'
    except:
        return get_notice_base_name()

def get_today_mmdd():
    today = tw_now()
    return f"{today.month:02d}{today.day:02d}"

# ── 計算已發送截止日 ──────────────────────────────────────────
def get_sent_deadline(today=None):
    if today is None:
        today = tw_now().replace(tzinfo=None)
    days_since_monday = today.weekday()
    this_monday = today - timedelta(days=days_since_monday)
    this_monday = this_monday.replace(hour=0, minute=0, second=0, microsecond=0)
    return this_monday + timedelta(days=20)

# ── 解析班表 ─────────────────────────────────────────────────
def load_schedule(file_obj):
    wb = openpyxl.load_workbook(file_obj)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hi = -1
    for i, row in enumerate(rows):
        vals = [str(v or '').strip() for v in row]
        if '店號' in vals and '日期' in vals:
            hi = i; break
    if hi < 0:
        raise ValueError('找不到欄位標題（店號/日期）')
    headers = [str(v or '').strip() for v in rows[hi]]
    idx_shop = headers.index('店號')
    idx_name = headers.index('店名')
    idx_date = headers.index('日期')
    idx_noon = headers.index('午別')
    map_ = {}
    first_date = None
    last_date = None
    for row in rows[hi+1:]:
        shop = str(row[idx_shop] or '').strip()
        if not shop: continue
        d = str(row[idx_date] or '').strip()
        n = str(row[idx_noon] or '').strip()
        noon_full = '上午' if n == '上' else ('下午' if n == '下' else n)
        if d:
            if first_date is None: first_date = d
            last_date = d
            map_.setdefault(shop, {'店名': str(row[idx_name] or '').strip(), 'dates': set(), 'noon': {}})
            map_[shop]['dates'].add(d)
            map_[shop]['noon'][d] = noon_full
    return map_, first_date, last_date

def parse_date(s):
    parts = s.split('/')
    return datetime(int(parts[0]), int(parts[1]), int(parts[2]))

def extract_version(filename):
    m = re.search(r'(\d{8})', filename)
    return m.group(1) if m else ''

def extract_date_from_filename(filename):
    m = re.search(r'(\d{8})', filename)
    if m:
        d = m.group(1)
        return f"{int(d[0:4])}/{int(d[4:6])}/{int(d[6:8])}"
    return ''

# ── 比對並分類 ────────────────────────────────────────────────
def compare_and_classify(m1, m2, deadline):
    異動通知 = []
    補1通知  = []
    for k in sorted(set(list(m1.keys()) + list(m2.keys()))):
        a, b = m1.get(k), m2.get(k)
        if not a or not b: continue
        s1 = '|'.join(d + '_' + a['noon'].get(d,'') for d in sorted(a['dates']))
        s2 = '|'.join(d + '_' + b['noon'].get(d,'') for d in sorted(b['dates']))
        if s1 != s2:
            orig = sorted(a['dates'])[0]
            chg  = sorted(b['dates'])[0]
            orig_dt = parse_date(orig)
            chg_dt  = parse_date(chg)
            diff = {
                '店號':   k,
                '店名':   a['店名'],
                '原訂日': orig,
                '原訂午': a['noon'].get(orig, ''),
                '異動日': chg,
                '異動午': b['noon'].get(chg, ''),
            }
            if orig_dt <= deadline:
                異動通知.append(diff)
            elif chg_dt <= deadline:
                補1通知.append(diff)
    異動通知.sort(key=lambda x: x['原訂日'])
    補1通知.sort(key=lambda x: x['原訂日'])
    return 異動通知, 補1通知

# ── 產出單份異動通知書 Excel ──────────────────────────────────
def generate_single_notice(diffs, version_full, notifier=''):
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    tmp.close()
    shutil.copy(TEMPLATE_PATH, tmp.name)
    wb = load_workbook(tmp.name)
    ws = wb['行程異動連絡單']
    DATE_FMT = 'M"月"D"日"'
    REASON = '因內部聯絡條，故行程異動'

    merges = [str(m) for m in list(ws.merged_cells.ranges)]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))

    for i, d in enumerate(diffs):
        r_d1 = 6 + i * 5
        r_d2 = 8 + i * 5
        ws.cell(row=r_d1, column=1).value = d['店號']
        c = ws.cell(row=r_d1, column=4)
        c.value = parse_date(d['原訂日']); c.number_format = DATE_FMT
        ws.cell(row=r_d1, column=5).value = d['原訂午']
        ws.cell(row=r_d1, column=6).value = REASON
        ws.cell(row=r_d1, column=8).value = notifier if notifier else None
        ws.cell(row=r_d1, column=9).value = None
        c2 = ws.cell(row=r_d2, column=4)
        c2.value = parse_date(d['異動日']); c2.number_format = DATE_FMT
        ws.cell(row=r_d2, column=5).value = d['異動午']
        ws.cell(row=r_d2, column=8).value = version_full

    for i in range(len(diffs), MAX_PER_SHEET):
        r_d1 = 6 + i * 5; r_d2 = 8 + i * 5
        for col in [1, 4, 5, 6, 8, 9]:
            ws.cell(row=r_d1, column=col).value = None
            ws.cell(row=r_d2, column=col).value = None

    for coord in merges:
        ws.merge_cells(coord)

    wb.save(tmp.name)
    return tmp.name

# ── 產出異動通知書（超過10店自動分份）────────────────────────
def generate_notice_files(diffs, version_full, notifier='', first_date=''):
    chunks = [diffs[i:i+MAX_PER_SHEET] for i in range(0, len(diffs), MAX_PER_SHEET)]
    files = []
    # 用班表第一筆日期產生檔名前綴，例如 (11604)盤點行程異動通知書-
    base_name = get_notice_filename(first_date) if first_date else get_notice_base_name()
    mmdd = get_today_mmdd()  # 操作當天日期 MMDD

    for idx, chunk in enumerate(chunks):
        path = generate_single_notice(chunk, version_full, notifier)
        if len(chunks) == 1:
            fname = f'{base_name}{mmdd}.xlsx'
        else:
            fname = f'{base_name}{mmdd}-{idx+1}.xlsx'
        files.append((fname, path))
    return files

# ── 產出補1通知書 PDF ────────────────────────────────────────
def make_bu1_pdf(store_name, chg_date, chg_noon, out_path):
    roc_year = str(chg_date.year - 1911)
    month    = str(chg_date.month)
    day      = str(chg_date.day)
    weekday  = WEEKDAY_MAP[chg_date.weekday()]

    doc = pymupdf.open(TEMPLATE_BU1_PATH)
    page = doc[0]
    blocks = page.get_text("dict")

    replacements = []
    for block in blocks['blocks']:
        if 'lines' not in block: continue
        for line in block['lines']:
            for span in line['spans']:
                if span['color'] == 255 and span['text'].strip():
                    replacements.append(span)

    for span in replacements:
        x, y = span['origin']
        old_text = span['text'].strip()
        size = span['size']

        if 150 < y < 170 and old_text not in ['114','5','一','上午','下午']:
            new_text = store_name
        elif 210 < y < 220:
            if x < 50:      new_text = roc_year
            elif x < 120:   new_text = month
            elif x < 175:   new_text = day
            elif x < 265:   new_text = weekday
            else:            new_text = chg_noon
        else:
            continue

        bbox = span['bbox']
        rect = pymupdf.Rect(bbox[0]-2, bbox[1]-1, bbox[2]+10, bbox[3]+1)
        page.draw_rect(rect, color=(1,1,1), fill=(1,1,1))

        orig_right = bbox[2]
        font = pymupdf.Font('china-t')
        text_width = font.text_length(new_text, fontsize=size)
        new_x = orig_right - text_width
        if new_x < x - 5:
            new_x = x
        page.insert_text((new_x, y), new_text, fontsize=size, color=(0,0,1), fontname='china-t')

    doc.save(out_path)

# ── API ───────────────────────────────────────────────────────
@app.route('/health', methods=['GET'])
def health():
    return jsonify({
        'status': 'ok',
        'server_time':      tw_now().strftime('%Y/%m/%d %H:%M'),
        'template_name':    get_template_name(),
        'shop_update_date': get_setting('SHOP_UPDATE_DATE', ''),
    })

@app.route('/admin/records', methods=['POST'])
def admin_records():
    data = request.get_json()
    password  = data.get('password', '')
    if password != get_admin_password():
        return jsonify({'error': '密碼錯誤'}), 401
    date_from = data.get('date_from', '')
    date_to   = data.get('date_to', '')
    dept      = data.get('dept', '')
    operator  = data.get('operator', '')
    doc_type  = data.get('doc_type', '')
    limit     = int(data.get('limit', 100))
    try:
        ws = get_sheet()
        if not ws:
            return jsonify({'error': 'Google Sheets 連線失敗'}), 500
        all_rows = ws.get_all_values()
        if not all_rows:
            return jsonify({'records': [], 'headers': []})
        headers = all_rows[0]
        records = all_rows[1:]

        def match(r):
            if len(r) < 16: return False
            if date_from and r[0][:10] < date_from: return False
            if date_to   and r[0][:10] > date_to:   return False
            if dept      and dept not in r[2]:       return False
            if operator  and operator not in r[1]:   return False
            if doc_type  and doc_type != r[15]:      return False
            return True

        filtered = [r for r in records if match(r)]
        filtered = list(reversed(filtered))[:limit]
        return jsonify({'headers': headers, 'records': filtered, 'total': len(filtered)})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/admin/search', methods=['POST'])
def admin_search():
    data = request.get_json()
    password = data.get('password', '')
    keyword  = data.get('keyword', '').strip()
    if password != get_admin_password():
        return jsonify({'error': '密碼錯誤'}), 401
    if not keyword:
        return jsonify({'error': '請輸入查詢關鍵字'}), 400
    try:
        ws = get_sheet()
        if not ws:
            return jsonify({'error': 'Google Sheets 連線失敗'}), 500
        all_rows = ws.get_all_values()
        if not all_rows:
            return jsonify({'records': [], 'headers': []})
        headers = all_rows[0]
        results = [r for r in all_rows[1:] if len(r) > 10 and (keyword in r[9] or keyword in r[10])]
        results = list(reversed(results))
        return jsonify({'headers': headers, 'records': results})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/compare', methods=['POST'])
def api_compare():
    if 'file1' not in request.files or 'file2' not in request.files:
        return jsonify({'error': '請上傳兩個班表檔案'}), 400
    f1 = request.files['file1']
    f2 = request.files['file2']
    force = request.form.get('force', 'false') == 'true'
    try:
        m1, first_date1, last_date1 = load_schedule(f1)
        m2, first_date2, last_date2 = load_schedule(f2)

        fn1 = extract_version(f1.filename)
        fn2 = extract_version(f2.filename)
        if fn1 and fn2:
            from datetime import datetime as dt_
            d1 = dt_.strptime(fn1, '%Y%m%d')
            d2 = dt_.strptime(fn2, '%Y%m%d')

            if d2 < d1:
                return jsonify({
                    'error': '⚠️ 上傳位置可能錯誤！版本二（異動後班表）的檔案日期早於版本一（原始班表），請確認是否上傳正確。',
                    'warn_swap': True
                }), 400

            diff_days = abs((d2 - d1).days)
            if diff_days > 7 and not force:
                return jsonify({
                    'warn_gap': True,
                    'gap_days': diff_days,
                    'message': f'⚠️ 兩個版本的班表日期相差 {diff_days} 天，是否確認繼續？'
                }), 200

        ver8   = extract_version(f2.filename)
        month  = str(int(first_date1.split('/')[1])).zfill(2)
        version_full = month + '-' + ver8
        deadline = get_sent_deadline()
        異動通知, 補1通知 = compare_and_classify(m1, m2, deadline)
        異動通知.sort(key=lambda x: x['原訂日'])
        補1通知.sort(key=lambda x: x['原訂日'])
        notice_files = (len(異動通知) + MAX_PER_SHEET - 1) // MAX_PER_SHEET if 異動通知 else 0

        operator = request.form.get('operator', '')
        dept     = request.form.get('dept', '')

        try:
            ws = get_sheet()
            log_to_sheet(ws, operator, dept, f1.filename, f2.filename, version_full, 異動通知, 補1通知)
        except Exception as e:
            print(f'Sheets log error: {e}')

        return jsonify({
            'version':       version_full,
            'month':         month,
            'v1_count':      len(m1),
            'v2_count':      len(m2),
            'deadline':      deadline.strftime('%Y/%m/%d'),
            'notice_count':  len(異動通知),
            'bu1_count':     len(補1通知),
            'notice_diffs':  異動通知,
            'bu1_diffs':     補1通知,
            'notice_files':  notice_files,
            'f1_date':       extract_date_from_filename(f1.filename),
            'f2_date':       extract_date_from_filename(f2.filename),
            'template_name':    get_template_name(),
            'shop_update_date': get_setting('SHOP_UPDATE_DATE', ''),
            'notice_filename':  get_notice_filename(first_date1),  # ← 新增：(11505)盤點行程異動通知書-
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/generate_notice', methods=['POST'])
def api_generate_notice():
    if 'file1' not in request.files or 'file2' not in request.files:
        return jsonify({'error': '請上傳兩個班表檔案'}), 400
    f1 = request.files['file1']
    f2 = request.files['file2']
    notifier = request.form.get('notifier', '')
    try:
        m1, first_date1, _ = load_schedule(f1)
        m2, _, _           = load_schedule(f2)
        print(f'[DEBUG] first_date1={first_date1}')
        print(f'[DEBUG] notice_filename={get_notice_filename(first_date1)}')
        ver8   = extract_version(f1.filename)
        month  = str(int(first_date1.split('/')[1])).zfill(2)
        version_full = month + '-' + ver8
        deadline = get_sent_deadline()
        異動通知, _ = compare_and_classify(m1, m2, deadline)
        if not 異動通知:
            return jsonify({'error': '無異動通知書資料'}), 400

        files = generate_notice_files(異動通知, version_full, notifier, first_date=first_date1)

        if len(files) == 1:
            return send_file(files[0][1], as_attachment=True,
                download_name=files[0][0],
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        else:
            zip_tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.zip')
            zip_tmp.close()
            with zipfile.ZipFile(zip_tmp.name, 'w') as zf:
                for fname, fpath in files:
                    zf.write(fpath, fname)
            base_name = get_notice_filename(first_date1) if first_date1 else get_notice_base_name()
            mmdd = get_today_mmdd()
            return send_file(zip_tmp.name, as_attachment=True,
                download_name=f'{base_name}{mmdd}.zip',
                mimetype='application/zip')
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/generate_bu1', methods=['POST'])
def api_generate_bu1():
    if 'file1' not in request.files or 'file2' not in request.files:
        return jsonify({'error': '請上傳兩個班表檔案'}), 400
    f1 = request.files['file1']
    f2 = request.files['file2']
    try:
        m1, first_date1, _ = load_schedule(f1)
        m2, _, _           = load_schedule(f2)
        ver8   = extract_version(f1.filename)
        month  = str(int(first_date1.split('/')[1])).zfill(2)
        version_full = month + '-' + ver8
        deadline = get_sent_deadline()
        _, 補1通知 = compare_and_classify(m1, m2, deadline)
        if not 補1通知:
            return jsonify({'error': '無補1通知書資料'}), 400

        tmpdir = tempfile.mkdtemp()
        pdf_files = []
        for d in 補1通知:
            chg_dt = parse_date(d['異動日'])
            mmdd = f"{chg_dt.month:02d}{chg_dt.day:02d}"
            store_display = d['店名'][:-1] if d['店名'].endswith('店') else d['店名']
            fname = f"(補1)實地盤點實施通知書-{store_display}{mmdd}.pdf"
            fpath = os.path.join(tmpdir, fname)
            make_bu1_pdf(d['店名'], chg_dt, d['異動午'], fpath)
            pdf_files.append((fname, fpath))

        mmdd_today = get_today_mmdd()
        zip_tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.zip')
        zip_tmp.close()
        with zipfile.ZipFile(zip_tmp.name, 'w') as zf:
            for fname, fpath in pdf_files:
                zf.write(fpath, fname)

        return send_file(zip_tmp.name, as_attachment=True,
            download_name=f'補1通知書_{version_full}_{mmdd_today}.zip',
            mimetype='application/zip')
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/admin/verify', methods=['POST'])
def admin_verify():
    data = request.get_json()
    password = data.get('password', '')
    admin_pw = get_setting('ADMIN_PASSWORD', os.environ.get('ADMIN_PASSWORD', ''))
    if password == admin_pw:
        return jsonify({
            'ok': True,
            'template_name':    get_setting('TEMPLATE_NAME', 'template.xlsx'),
            'shop_update_date': get_setting('SHOP_UPDATE_DATE', ''),
            'staff_update_date':get_setting('STAFF_UPDATE_DATE', ''),
        })
    return jsonify({'ok': False, 'error': '密碼錯誤'}), 401

@app.route('/admin/change_password', methods=['POST'])
def admin_change_password():
    data = request.get_json()
    old_pw  = data.get('old_password', '')
    new_pw  = data.get('new_password', '')
    if old_pw != get_admin_password():
        return jsonify({'error': '舊密碼錯誤'}), 401
    if not new_pw or len(new_pw) < 6:
        return jsonify({'error': '新密碼至少需要 6 個字元'}), 400

    ok = set_setting('ADMIN_PASSWORD', new_pw)
    if ok:
        return jsonify({'ok': True, 'message': '密碼已更新！立即生效。'})
    else:
        return jsonify({'error': '更新失敗，請稍後再試'}), 500

@app.route('/admin/upload_template', methods=['POST'])
def admin_upload_template():
    password = request.form.get('password', '')
    if password != get_admin_password():
        return jsonify({'error': '密碼錯誤'}), 401
    if 'template' not in request.files:
        return jsonify({'error': '請上傳範本檔案'}), 400
    f = request.files['template']
    if not f.filename.endswith('.xlsx'):
        return jsonify({'error': '請上傳 .xlsx 格式'}), 400
    f.save(TEMPLATE_PATH)
    save_template_name(f.filename)
    upload_template_to_drive()  # ✅ 同步到 Drive
    return jsonify({
        'ok': True,
        'message': f'範本已更新！({f.filename})',
        'template_name': f.filename,
        'shop_update_date':  get_setting('SHOP_UPDATE_DATE', ''),
        'staff_update_date': get_setting('STAFF_UPDATE_DATE', ''),
    })

@app.route('/admin/upload_shop', methods=['POST'])
def admin_upload_shop():
    """上傳店舖資料，更新範本裡的「店舖」工作表"""
    password = request.form.get('password', '')
    if password != get_admin_password():
        return jsonify({'error': '密碼錯誤'}), 401
    if 'file' not in request.files:
        return jsonify({'error': '請上傳店舖檔案'}), 400
    f = request.files['file']
    if not f.filename.endswith('.xlsx'):
        return jsonify({'error': '請上傳 .xlsx 格式'}), 400
    try:
        # 讀取上傳的店舖資料
        new_wb = load_workbook(f)
        new_ws = new_wb.active

        # 讀取現有範本
        if not os.path.exists(TEMPLATE_PATH):
            return jsonify({'error': '範本檔案不存在，請先上傳完整範本'}), 400
        tmpl_wb = load_workbook(TEMPLATE_PATH)

        # 確認範本有「店舖」工作表
        if '店舖' not in tmpl_wb.sheetnames:
            return jsonify({'error': '範本裡找不到「店舖」工作表'}), 400

        # 清空並重寫「店舖」工作表
        tmpl_ws = tmpl_wb['店舖']
        for row in tmpl_ws.iter_rows():
            for cell in row:
                cell.value = None

        for i, row in enumerate(new_ws.iter_rows(values_only=True)):
            for j, val in enumerate(row):
                tmpl_ws.cell(row=i+1, column=j+1, value=val)

        tmpl_wb.save(TEMPLATE_PATH)

        # ✅ 同步到 Drive
        upload_template_to_drive()

        # 更新 Sheets 紀錄日期
        today_str = tw_now().strftime('%Y/%m/%d')
        set_setting('SHOP_UPDATE_DATE', today_str)

        return jsonify({
            'ok': True,
            'message': f'店舖資料已更新！({today_str})',
            'shop_update_date': today_str,
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/admin/upload_staff', methods=['POST'])
def admin_upload_staff():
    """上傳擔當資料，更新範本裡的「擔當公務機」工作表"""
    password = request.form.get('password', '')
    if password != get_admin_password():
        return jsonify({'error': '密碼錯誤'}), 401
    if 'file' not in request.files:
        return jsonify({'error': '請上傳擔當檔案'}), 400
    f = request.files['file']
    if not f.filename.endswith('.xlsx'):
        return jsonify({'error': '請上傳 .xlsx 格式'}), 400
    try:
        new_wb = load_workbook(f)
        new_ws = new_wb.active

        if not os.path.exists(TEMPLATE_PATH):
            return jsonify({'error': '範本檔案不存在，請先上傳完整範本'}), 400
        tmpl_wb = load_workbook(TEMPLATE_PATH)

        if '擔當公務機' not in tmpl_wb.sheetnames:
            return jsonify({'error': '範本裡找不到「擔當公務機」工作表'}), 400

        tmpl_ws = tmpl_wb['擔當公務機']
        for row in tmpl_ws.iter_rows():
            for cell in row:
                cell.value = None

        for i, row in enumerate(new_ws.iter_rows(values_only=True)):
            for j, val in enumerate(row):
                tmpl_ws.cell(row=i+1, column=j+1, value=val)

        tmpl_wb.save(TEMPLATE_PATH)

        # ✅ 同步到 Drive
        upload_template_to_drive()

        today_str = tw_now().strftime('%Y/%m/%d')
        set_setting('STAFF_UPDATE_DATE', today_str)

        return jsonify({
            'ok': True,
            'message': f'擔當資料已更新！({today_str})',
            'staff_update_date': today_str,
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
